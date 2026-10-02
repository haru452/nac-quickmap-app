"""Excel (.xlsx) 出力。Qt には依存しない。"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEADERS: tuple[str, ...] = (
    "Product ID", "Original Product", "Target", "Orbit",
    "Center Latitude", "Center Longitude", "Resolution",
    "Scaled Pixel Width", "Scaled Pixel Height", "Slew Angle",
    "Incidence Angle", "Emission Angle", "Phase Angle", "Source URL",
    "Registered At",
)

KEYS: tuple[str, ...] = (
    "product_id", "original_product", "target_name", "orbit_number",
    "center_latitude", "center_longitude", "resolution",
    "scaled_pixel_width", "scaled_pixel_height", "slew_angle",
    "incidence_angle", "emission_angle", "phase_angle", "source_url",
    "created_at",
)

WIDTHS: tuple[int, ...] = (16, 18, 10, 8, 16, 16, 12, 16, 16, 12, 16, 16, 12, 60, 20)


def export_xlsx(records: Sequence[Mapping[str, Any]], path: str | Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LROC"
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for record in records:
        sheet.append([_cell(record, key) for key in KEYS])
    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    workbook.save(str(path))


def _cell(record: Mapping[str, Any], key: str) -> Any:
    value = record.get(key)
    return "" if value is None else value
