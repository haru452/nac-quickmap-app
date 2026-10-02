from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl

from app import excel

RECORD = {
    "product_id": "M190737496RE",
    "original_product": "nacr00098809",
    "target_name": "MOON",
    "orbit_number": 13151,
    "center_latitude": -0.3,
    "center_longitude": 339.1,
    "resolution": 0.9408832674492926,
    "scaled_pixel_width": 1.1,
    "scaled_pixel_height": 0.78,
    "slew_angle": 0.012741557399709086,
    "incidence_angle": 45.78,
    "emission_angle": 1.18,
    "phase_angle": 44.6,
    "source_url": "https://data.lroc.im-ldi.com/lroc/view_lroc/LRO-L-LROC-3-CDR-V1.0/M190737496RC",
    "created_at": "2026-10-01 12:00:00",
}


class ExportXlsxTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.out = Path(self.tmp.name) / "out.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def test_header_order_matches_spec(self):
        self.assertEqual(
            list(excel.HEADERS),
            ["Product ID", "Original Product", "Target", "Orbit",
             "Center Latitude", "Center Longitude", "Resolution",
             "Scaled Pixel Width", "Scaled Pixel Height", "Slew Angle",
             "Incidence Angle", "Emission Angle", "Phase Angle", "Source URL", "Registered At"])
        self.assertEqual(len(excel.KEYS), len(excel.HEADERS))

    def test_keys_order_matches_spec(self):
        self.assertEqual(
            excel.KEYS,
            ("product_id", "original_product", "target_name", "orbit_number",
             "center_latitude", "center_longitude", "resolution",
             "scaled_pixel_width", "scaled_pixel_height", "slew_angle",
             "incidence_angle", "emission_angle", "phase_angle", "source_url", "created_at"))

    def test_header_and_numeric_cells(self):
        excel.export_xlsx([RECORD], self.out)
        sheet = openpyxl.load_workbook(self.out).active
        self.assertEqual(sheet.cell(1, 1).value, "Product ID")
        self.assertEqual(sheet.cell(1, 1).font.bold, True)
        self.assertEqual(sheet.cell(2, 1).value, "M190737496RE")
        self.assertIsInstance(sheet.cell(2, 4).value, int)
        self.assertEqual(sheet.cell(2, 4).value, 13151)
        self.assertIsInstance(sheet.cell(2, 7).value, float)
        self.assertAlmostEqual(sheet.cell(2, 8).value, 1.1)
        self.assertAlmostEqual(sheet.cell(2, 9).value, 0.78)
        self.assertAlmostEqual(sheet.cell(2, 10).value, 0.012741557399709086)
        self.assertEqual(sheet.freeze_panes, "A2")

    def test_missing_values_become_blank(self):
        excel.export_xlsx([{"product_id": "X"}], self.out)
        sheet = openpyxl.load_workbook(self.out).active
        self.assertIn(sheet.cell(2, 10).value, ("", None))

    def test_no_records_writes_header_only(self):
        excel.export_xlsx([], self.out)
        sheet = openpyxl.load_workbook(self.out).active
        self.assertEqual(sheet.max_row, 1)

    def test_unwritable_path_raises_oserror(self):
        with self.assertRaises(OSError):
            excel.export_xlsx([RECORD], Path(self.tmp.name) / "missing" / "dir.xlsx")
